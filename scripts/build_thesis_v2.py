"""Build the revised chapters and a separate editable evidence appendix.

Uses the same Word-template helpers as version one. Existing version-one files
are never changed. The new evidence script must run before this builder.
"""

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import hashlib
import json
import re
import shutil

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import build_thesis_docx as word

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/thesis_v2"
EVIDENCE = REPORT / "evidence"


def add_table(doc, spec):
    word.caption(doc, spec["title"], spec["note"])
    t = word.table(doc, spec["headers"], spec["rows"], spec["widths"])
    for i, row in enumerate(t.rows):
        for j, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if j in spec["left"]
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                # These compact report tables fit on a page; keep each intact.
                p.paragraph_format.keep_with_next = i < len(t.rows) - 1
    return t


def build(source, destination, tables):
    word.FOOTNOTES.clear()
    before = hashlib.sha256(word.TEMPLATE.read_bytes()).hexdigest()
    doc = Document(word.TEMPLATE)
    section = deepcopy(doc.sections[-1]._sectPr)
    for child in list(doc._element.body):
        doc._element.body.remove(child)
    doc._element.body.append(section)
    for tag in ("headerReference", "footerReference", "titlePg"):
        for element in section.findall(qn("w:" + tag)):
            section.remove(element)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[name]
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap = doc.styles.add_style("Thesis Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap.base_style = normal
    cap.font.size = Pt(10)
    cap.paragraph_format.line_spacing = 1.1
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(5)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    reference_mode = False
    pagebreak = False
    for block in re.split(r"\n\s*\n", source.read_text().strip()):
        block = block.strip()
        if block.startswith("<!-- PAGEBREAK -->"):
            pagebreak = True
            block = block.replace("<!-- PAGEBREAK -->", "").strip()
            if not block:
                continue
        if block.startswith("#"):
            level = len(block) - len(block.lstrip("#"))
            title = block[level:].strip()
            p = doc.add_paragraph(title, style=f"Heading {level}")
            word.explicit_numbering_off(p)
            if pagebreak:
                p.paragraph_format.page_break_before = True
                pagebreak = False
            reference_mode = title == "References"
        elif block.startswith("{{TABLE:"):
            add_table(doc, tables[block[8:-2]])
        elif block.startswith("{{EQUATION:"):
            word.add_equation(doc, block[11:-2])
        elif block.startswith("{{FIGURE:"):
            name, title, note = block[9:-2].split("|")
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(
                str(EVIDENCE / "figures" / f"{name}.png"), width=Cm(16)
            )
            word.caption(doc, title, note, keep=False)
        else:
            p = word.add_text(doc.add_paragraph(), block)
            if reference_mode:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together = True
                for r in p.runs:
                    r.font.size = Pt(10)
    word.attach_footnotes(doc)
    doc.core_properties.title = (
        "NLP and machine learning analysis of private equity literature"
    )
    doc.core_properties.subject = "Methods and thematic results on 292 distinct works"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    if destination.exists():
        backup = ROOT.parent / "batill_thesis_draft_archive" / "v2_revisions"
        backup.mkdir(parents=True, exist_ok=True)
        shutil.copy2(
            destination,
            backup / f"{destination.stem}_{datetime.now():%Y%m%d_%H%M%S}.docx",
        )
    doc.save(destination)
    assert hashlib.sha256(word.TEMPLATE.read_bytes()).hexdigest() == before
    print(destination)


def appendix_source():
    sections = [
        (
            "A Model selection and sensitivity",
            "Full resolution sweeps, seeds and subset memberships are retained in the evidence folder. These diagnostics are not label-accuracy estimates.",
            ["seed_summary", "full_sweep"],
        ),
        (
            "B Alternative thematic granularity",
            "The following K-means profiles show how additional groups change thematic detail and separation. All corresponding text-Leiden profiles are in labelled_profiles.csv and cluster_review_cards.md. Names follow lexical and paper inspection, not a fitted classifier.",
            [f"profiles{k}" for k in range(4, 7)],
        ),
    ]
    for c in range(1, 4):
        sections.append(
            (
                f"C{c} Expression evidence for main cluster C{c}",
                "The pooled ranking is class-based TF-IDF; contrast compares normalized per-paper TF-IDF inside and outside the group. Both tables use the same vocabulary and minimum-support rule. Similar expressions are retained rather than manually consolidated.",
                [f"terms{c}_contrast", f"terms{c}_pooled"],
            )
        )
    for c in range(1, 4):
        sections.append(
            (
                f"D{c} Paper evidence for main cluster C{c}",
                "Central papers, lowest-silhouette boundary cases and seeded random examples expose both typical and difficult assignments. Source excerpts and their page/block references are retained in review_papers.csv. A negative silhouette is a review signal, not a reason for automatic exclusion.",
                [f"examples{c}"],
            )
        )
    for method in ("kmeans", "leiden"):
        sections.append(
            (
                f'E {"K means" if method=="kmeans" else "Text Leiden"} changes between counts',
                "Each cell counts papers shared by a group before and after changing granularity. The partitions are independently fitted and need not be nested. Theme names and paper lists are in the corresponding profiles.",
                [f"transition_{method}_{k}" for k in (3, 4, 5)],
            )
        )
    sections += [
        (
            "F Citation evidence and coverage sensitivity",
            "The citation baseline uses resolution one. Its representatives use within-community degree. Dropping manually verified edges changes coverage as well as membership; comparisons use only works connected in both versions.",
            ["citation_sensitivity", "citation_papers"],
        ),
        (
            "G Hierarchical inspection",
            "The 15-group cut is an inspection aid. Cluster size alone does not establish irrelevance or a distinct research theme. All member titles are saved in hierarchy_membership_15.csv.",
            ["hierarchy_counts"],
        ),
    ]
    blocks = [
        "# Supporting evidence for the literature analysis",
        "The claim–evidence index links these tables to saved data and source-linked examples. Main-map cluster IDs refer to K-means at three clusters unless specified otherwise.",
    ]
    for index, (title, intro, keys) in enumerate(sections):
        if index:
            blocks.append("<!-- PAGEBREAK -->")
        blocks.extend(["## " + title, intro])
        blocks.extend("{{TABLE:" + key + "}}" for key in keys)
    path = REPORT / "supporting_appendix.md"
    path.write_text("\n\n".join(blocks) + "\n")
    return path


if __name__ == "__main__":
    tables = json.loads((EVIDENCE / "tables.json").read_text())
    build(REPORT / "thesis_sections.md", REPORT / "thesis_sections.docx", tables)
    build(appendix_source(), REPORT / "supporting_appendix.docx", tables)

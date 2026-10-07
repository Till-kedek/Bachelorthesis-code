"""Build an editable chapter insert using the supplied Word template.

Edit reports/thesis/thesis_sections.md for reproducible prose changes. Direct Word
edits are also supported: this builder backs up an existing output before writing.
The source template is never modified. Figures and numeric tables come from the
frozen analysis evidence. No analysis or external requests run during this build.
"""

from copy import deepcopy
from datetime import datetime
from pathlib import Path
import argparse
import hashlib
import re
import shutil

import pandas as pd
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Pt, Cm, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / "reports/thesis"
EVIDENCE = REPORT / "evidence"
TEMPLATE = ROOT / "Thesis_Template_FinancialManagement_202502.docx"
FOOTNOTES = []


def explicit_numbering_off(p):
    props = p._p.get_or_add_pPr()
    num = OxmlElement("w:numPr")
    ident = OxmlElement("w:numId")
    ident.set(qn("w:val"), "0")
    num.append(ident)
    props.append(num)


def add_text(p, text):
    for part in re.split(r"(`[^`]+`)", text):
        if part.startswith("`"):
            run = p.add_run(part.strip("`"))
            run.font.name = "Courier New"
            run.font.size = Pt(8.5)
        else:
            for piece in re.split(r"(\b[zxwdc]_[a-z]+\b)", part):
                if re.fullmatch(r"[zxwdc]_[a-z]+", piece):
                    base, index = piece.split("_")
                    p.add_run(base).italic = True
                    r = p.add_run(index)
                    r.italic = True
                    r.font.subscript = True
                else:
                    p.add_run(piece)
    return p


def footnote(p, text):
    FOOTNOTES.append(text)
    r = p.add_run()
    mark = OxmlElement("w:footnoteReference")
    mark.set(qn("w:id"), str(len(FOOTNOTES)))
    r._r.append(mark)
    r.font.superscript = True


def caption(doc, text, note="", keep=True):
    p = doc.add_paragraph(style="Thesis Caption")
    p.paragraph_format.keep_with_next = keep
    p.paragraph_format.keep_together = True
    p.add_run(text).bold = True
    if note:
        p.add_run("\n" + note)
    footnote(
        p,
        "Own source. Analysis of the final 292-work corpus.",
    )
    return p


def table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for cell, w in zip(t.columns, widths):
        cell.width = Cm(w)
    for c, v in zip(t.rows[0].cells, headers):
        c.text = str(v)
    repeat = OxmlElement("w:tblHeader")
    t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells = t.add_row().cells
        for c, v in zip(cells, row):
            c.text = str(v)
    for ri, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        for j, c in enumerate(row.cells):
            c.width = Cm(widths[j])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = c._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
                b = OxmlElement("w:" + edge)
                b.set(qn("w:val"), "single")
                b.set(qn("w:sz"), "4")
                b.set(qn("w:color"), "D9D9D9")
                borders.append(b)
            tcpr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in [
                ("top", 65),
                ("bottom", 65),
                ("left", 90),
                ("right", 90),
            ]:
                m = OxmlElement("w:" + edge)
                m.set(qn("w:w"), str(value))
                m.set(qn("w:type"), "dxa")
                margins.append(m)
            tcpr.append(margins)
            if ri == 0:
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E6E6E6")
                tcpr.append(shade)
            for p in c.paragraphs:
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.keep_with_next = ri < len(t.rows) - 1
                p.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                )
                for r in p.runs:
                    r.font.name = "Arial"
                    r.font.size = Pt(9.5)
                    r.bold = ri == 0
    return t


def add_table(doc, name):
    if name == "corpus_flow":
        caption(
            doc,
            "Table 1 Corpus construction",
            "Counts refer to files until canonical works are established.",
        )
        rows = [
            ["Supplied PDF files", "320", "Starting folder"],
            ["Unique byte-level PDFs", "319", "1 exact duplicate removed"],
            ["Usable text and stored embeddings", "318", "1 empty extraction excluded"],
            ["In-scope selected PDFs", "294", "24 reviewed scope exclusions"],
            [
                "Canonical works for analysis",
                "292",
                "2 duplicate-version pairs collapsed",
            ],
        ]
        table(doc, ["Stage", "Number", "Change"], rows, [8.5, 1.7, 5.8])
    elif name == "selection":
        caption(
            doc,
            "Table 2 Separation and stability for the reviewed cluster counts",
            "Sil. = mean cosine silhouette. Seed = mean ARI across initialisation seeds. Sub. = mean ARI on shared works in five 80% K-means subsamples. Cross = ARI between text methods. Leiden settings are selected from the full sweep.",
        )
        m = pd.read_csv(EVIDENCE / "main_selection_table.csv")
        rows = []
        for _, r in m.iterrows():
            rows.append(
                [
                    int(r.k),
                    f"{r.kmeans_silhouette:.3f}",
                    f"{r.leiden_silhouette:.3f}",
                    f"{r.kmeans_seed_ARI:.3f}",
                    f"{r.leiden_seed_ARI:.3f}",
                    f"{r.kmeans_subsample_ARI:.3f}",
                    f"{r.cross_method_ARI:.3f}",
                ]
            )
        table(
            doc,
            [
                "k",
                "K-means\nSil.",
                "Leiden\nSil.",
                "K-means\nSeed",
                "Leiden\nSeed",
                "K-means\nSub.",
                "Cross",
            ],
            rows,
            [1, 2.5, 2.5, 2.5, 2.5, 2.5, 1.5],
        )
    elif name == "themes":
        caption(
            doc,
            "Table 3 Provisional themes in the six-cluster text fits",
            "Cluster numbers refer to the displayed fits. Numbers are works; themes are qualitative interpretations of terms and reviewed examples.",
        )
        labels = pd.read_csv(EVIDENCE / "labelled_profiles.csv")
        rows = []
        names = [
            "Sponsor practices and portfolio governance",
            "Buyout financing and transaction value",
            "Public versus private financing and listing",
            "Employment and operating outcomes",
            "Fund investment and performance",
            "Healthcare ownership and outcomes",
        ]
        for c, name in enumerate(names, 1):
            a = labels[(labels.solution == "kmeans_k6") & (labels.cluster == c)].iloc[0]
            b = labels[(labels.solution == "leiden_k6") & (labels.cluster == c)].iloc[0]
            rows.append([f"{c}  {name}", int(a.papers), int(b.papers)])
        table(doc, ["Dominant theme", "K-means", "Text Leiden"], rows, [11.4, 2.3, 2.3])
    else:
        raise ValueError(name)


def add_equation(doc, name):
    """Native OMML with structural subscripts, avoiding Unicode pseudo-scripts."""

    def run(text):
        r = OxmlElement("m:r")
        pr = OxmlElement("m:rPr")
        style = OxmlElement("m:sty")
        style.set(qn("m:val"), "p")
        pr.append(style)
        r.append(pr)
        t = OxmlElement("m:t")
        t.text = text
        r.append(t)
        return r

    def sub(base, index):
        obj = OxmlElement("m:sSub")
        e = OxmlElement("m:e")
        e.append(run(base))
        obj.append(e)
        child = OxmlElement("m:sub")
        child.append(run(index))
        obj.append(child)
        return obj

    def norm():
        obj = OxmlElement("m:sSub")
        e = OxmlElement("m:e")
        e.extend([run("‖"), sub("u", "p"), run("‖")])
        obj.append(e)
        child = OxmlElement("m:sub")
        child.append(run("2"))
        obj.append(child)
        return obj

    if name == "paper_vector":
        parts = [
            sub("u", "p"),
            run(" = "),
            sub("∑", "s"),
            run(" "),
            sub("w", "ps"),
            run(" "),
            sub("z", "ps"),
            run(" ,     "),
            sub("x", "p"),
            run(" = "),
            sub("u", "p"),
            run(" / "),
            norm(),
        ]
    elif name == "leiden":
        parts = [
            run("Q(γ) = "),
            sub("∑", "ij"),
            run(" [ "),
            sub("A", "ij"),
            run(" − γ "),
            sub("d", "i"),
            run(" "),
            sub("d", "j"),
            run(" / (2m) ]  1("),
            sub("c", "i"),
            run(" = "),
            sub("c", "j"),
            run(")"),
        ]
    else:
        raise ValueError(name)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.15
    math = OxmlElement("m:oMath")
    math.extend(parts)
    p._p.append(math)


def attach_footnotes(doc):
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    root = etree.Element(qn("w:footnotes"), nsmap={"w": ns})
    for ident, kind in [(-1, "separator"), (0, "continuationSeparator")]:
        fn = etree.SubElement(
            root, qn("w:footnote"), {qn("w:type"): kind, qn("w:id"): str(ident)}
        )
        p = etree.SubElement(fn, qn("w:p"))
        r = etree.SubElement(p, qn("w:r"))
        etree.SubElement(r, qn("w:" + kind))
    for ident, text in enumerate(FOOTNOTES, 1):
        fn = etree.SubElement(root, qn("w:footnote"), {qn("w:id"): str(ident)})
        p = etree.SubElement(fn, qn("w:p"))
        pr = etree.SubElement(p, qn("w:pPr"))
        etree.SubElement(
            pr,
            qn("w:spacing"),
            {
                qn("w:line"): "240",
                qn("w:lineRule"): "auto",
                qn("w:before"): "0",
                qn("w:after"): "0",
            },
        )
        for isref, txt in [(True, ""), (False, " " + text)]:
            r = etree.SubElement(p, qn("w:r"))
            rp = etree.SubElement(r, qn("w:rPr"))
            etree.SubElement(
                rp, qn("w:rFonts"), {qn("w:ascii"): "Arial", qn("w:hAnsi"): "Arial"}
            )
            etree.SubElement(rp, qn("w:sz"), {qn("w:val"): "18"})
            if isref:
                etree.SubElement(rp, qn("w:vertAlign"), {qn("w:val"): "superscript"})
                etree.SubElement(r, qn("w:footnoteRef"))
            else:
                el = etree.SubElement(r, qn("w:t"))
                el.text = txt
                el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    for rid, rel in list(doc.part.rels.items()):
        if rel.reltype == RT.FOOTNOTES:
            del doc.part.rels[rid]
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True),
        doc.part.package,
    )
    doc.part.relate_to(part, RT.FOOTNOTES)


def build(destination=None):
    FOOTNOTES.clear()
    destination = Path(destination or REPORT / "thesis_sections.docx")
    before = hashlib.sha256(TEMPLATE.read_bytes()).hexdigest()
    doc = Document(TEMPLATE)
    section = deepcopy(doc.sections[-1]._sectPr)
    for child in list(doc._element.body):
        doc._element.body.remove(child)
    doc._element.body.append(section)
    for el in section.findall(qn("w:headerReference")) + section.findall(
        qn("w:footerReference")
    ):
        section.remove(el)
    for el in section.findall(qn("w:titlePg")):
        section.remove(el)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(6)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for name in ["Heading 1", "Heading 2", "Heading 3"]:
        s = doc.styles[name]
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.line_spacing = 1.15
        s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    cap = doc.styles.add_style("Thesis Caption", WD_STYLE_TYPE.PARAGRAPH)
    cap.base_style = normal
    cap.font.size = Pt(10)
    cap.paragraph_format.line_spacing = 1.1
    cap.paragraph_format.space_before = Pt(6)
    cap.paragraph_format.space_after = Pt(5)
    cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    sec = doc.sections[0]
    sec.different_first_page_header_footer = False
    doc.settings.odd_and_even_pages_header_footer = False
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.paragraph_format.space_before = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    field.append(r)
    f._p.append(field)
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)
    reference_mode = False
    pagebreak = False
    for block in re.split(
        r"\n\s*\n", (REPORT / "thesis_sections.md").read_text().strip()
    ):
        block = block.strip()
        if not block:
            continue
        if block.startswith("<!-- PAGEBREAK -->"):
            pagebreak = True
            block = block.replace("<!-- PAGEBREAK -->", "").strip()
            if not block:
                continue
        if block.startswith("#"):
            level = len(block) - len(block.lstrip("#"))
            text = block[level:].strip()
            p = doc.add_paragraph(text, style=f"Heading {level}")
            explicit_numbering_off(p)
            if pagebreak:
                p.paragraph_format.page_break_before = True
                pagebreak = False
            if text == "References":
                reference_mode = True
        elif block.startswith("{{TABLE:"):
            add_table(doc, block[8:-2])
        elif block.startswith("{{EQUATION:"):
            add_equation(doc, block[11:-2])
        elif block.startswith("{{FIGURE:"):
            name, title, note = block[9:-2].split("|")
            p = doc.add_paragraph()
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.line_spacing = 1
            width = 16 if name == "umap_comparison" else 11
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(
                str(EVIDENCE / "figures" / f"{name}.png"), width=Cm(width)
            )
            caption(doc, title, note, keep=False)
        else:
            p = add_text(doc.add_paragraph(), block)
            if reference_mode:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together = True
                for r in p.runs:
                    r.font.size = Pt(10)
    attach_footnotes(doc)
    doc.core_properties.title = (
        "Methods results and technical discussion for a map of private equity research"
    )
    doc.core_properties.subject = (
        "Data-science chapter insert based on 292 canonical works"
    )
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    if destination.exists():
        backup = REPORT / "document_backups"
        backup.mkdir(exist_ok=True)
        shutil.copy2(
            destination, backup / f"thesis_sections_{datetime.now():%Y%m%d_%H%M%S}.docx"
        )
    doc.save(destination)
    assert hashlib.sha256(TEMPLATE.read_bytes()).hexdigest() == before
    print(destination)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    build(args.output)

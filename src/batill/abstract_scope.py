"""Auditable subject screening and exact reuse of unchanged abstract embeddings.

This is a reproducible lexical screen, not a claim of manual relevance assessment.
Ambiguous records are withheld pending review. The original corpus is immutable.
"""

import hashlib
import json
from pathlib import Path
import re
import shutil
import tempfile
import unicodedata

import numpy as np
import pandas as pd

from .abstract_clustering import load_abstract_embeddings
from .storage import file_hash


ORIGINAL_SOURCE = 'data/abstracts/abstracts_clean.csv'
ORIGINAL_RUN = 'outputs/abstract_embeddings/bge_m3/20260930T174845_508054Z_full_3c498a9d'
ACTIVE_CONFIG = 'configs/abstract_analysis_input.json'
OVERRIDE_COLUMNS = ['paper_id', 'record_sha256', 'decision', 'reason']


def _json_hash(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def _write_json(path, value):
    Path(path).write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def _normalize(text):
    return re.sub(r'\s+', ' ', unicodedata.normalize('NFKC', str(text)).translate(
        str.maketrans({'–': '-', '—': '-', '‑': '-', '‐': '-', '\u00ad': ''}))).strip().casefold()


def screen_abstracts(source, rules, overrides=None):
    """Return a complete audit; retain only clear PE records or explicit overrides.

    Overrides require a reason and a hash of the complete original metadata row,
    so stale decisions cannot follow an EID whose text or metadata has changed.
    """
    source = source.copy().fillna('').astype(str)
    required = {'paper_id', 'title', 'abstract', 'embedding_text', 'journal', 'document_type'}
    if not required <= set(source):
        raise ValueError(f'Missing scope columns: {sorted(required - set(source))}')
    if source.empty or not source.paper_id.is_unique or source.paper_id.str.strip().eq('').any():
        raise ValueError('Scope source needs unique nonempty paper IDs and at least one record.')
    if rules.get('schema_version') != 1:
        raise ValueError('Unsupported scope rule schema.')
    if type(rules.get('exclude_any_vc_mention')) is not bool:
        raise ValueError('exclude_any_vc_mention must be a boolean.')
    minimum = rules['minimum_abstract_pe_mentions_without_title']
    if type(minimum) is not int or minimum < 1:
        raise ValueError('Minimum PE mentions must be a positive integer.')
    patterns = {name: re.compile(rules[name + '_pattern'], re.I)
                for name in ('private_equity', 'venture_capital', 'startup', 'chemical', 'ambiguous_financing')}
    trade = {_normalize(journal) for journal in rules['trade_journals']}
    research_types = {_normalize(kind) for kind in rules['research_document_types']}
    rows = []
    for row in source.to_dict('records'):
        title, abstract, journal = (_normalize(row[field]) for field in ('title', 'abstract', 'journal'))
        text = title + '\n' + abstract
        hits = {name: list(dict.fromkeys(m.group(0) for m in pattern.finditer(text)))
                for name, pattern in patterns.items()}
        pe_title = bool(patterns['private_equity'].search(title))
        pe_count = len(list(patterns['private_equity'].finditer(abstract)))
        chemical_source = bool(patterns['chemical'].search(journal))
        core = pe_title or pe_count >= minimum
        # Explicit exclusions precede positive evidence: a PE mention cannot rescue VC/news.
        if hits['chemical'] or chemical_source:
            decision, reason = 'exclude', 'chemical_industry_or_chemistry'
        elif journal in trade:
            decision, reason = 'exclude', 'industry_trade_news'
        elif hits['venture_capital'] or journal == 'venture capital':
            if rules['exclude_any_vc_mention'] or not core or journal == 'venture capital':
                decision, reason = 'exclude', 'venture_capital_or_mixed_pe_vc'
            else:
                decision, reason = 'review', 'mixed_pe_vc_requires_review'
        elif hits['startup']:
            decision, reason = 'review' if core else 'exclude', 'startup_or_early_stage_focus'
        elif hits['ambiguous_financing']:
            decision, reason = 'review', 'private_placement_or_crowdfunding_requires_review'
        elif not hits['private_equity']:
            decision, reason = 'exclude', 'no_explicit_private_equity_evidence'
        elif _normalize(row['document_type']) not in research_types:
            decision, reason = 'review', 'nonresearch_document_type'
        elif not core:
            decision, reason = 'review', 'incidental_or_uncertain_private_equity_focus'
        else:
            decision, reason = 'keep', 'explicit_private_equity_or_buyout_focus'
        rows.append({**row, 'record_sha256': _json_hash(row),
                     'automatic_decision': decision, 'automatic_reason': reason,
                     'pe_in_title': pe_title, 'pe_abstract_mentions': pe_count,
                     **{name + '_matches': '; '.join(values) for name, values in hits.items()},
                     'chemical_source': chemical_source, 'trade_source': journal in trade,
                     'decision': decision, 'reason': reason, 'manual_override': False})
    audit = pd.DataFrame(rows)
    if overrides is not None and not overrides.empty:
        overrides = overrides.fillna('').astype(str)
        if not set(OVERRIDE_COLUMNS) <= set(overrides):
            raise ValueError(f'Overrides require {OVERRIDE_COLUMNS}')
        if not overrides.paper_id.is_unique:
            raise ValueError('Duplicate override paper IDs.')
        by_id = audit.set_index('paper_id', drop=False)
        for entry in overrides.to_dict('records'):
            key = entry['paper_id']
            if key not in by_id.index:
                raise ValueError(f'Unknown override paper ID: {key}')
            if entry['record_sha256'] != by_id.at[key, 'record_sha256']:
                raise ValueError(f'Stale override record hash for {key}')
            if entry['decision'] not in {'keep', 'exclude', 'review'} or not entry['reason'].strip():
                raise ValueError(f'Override needs keep/exclude/review and a nonempty reason: {key}')
            by_id.loc[key, ['decision', 'reason', 'manual_override']] = [
                entry['decision'], entry['reason'].strip(), True]
        audit = by_id.reset_index(drop=True)
    return audit


def export_review_workbook(audit, destination):
    """Export an Excel view. Durable decisions live in the overrides CSV."""
    from openpyxl.styles import Alignment, Font, PatternFill
    counts = audit.groupby(['decision', 'reason']).size().rename('records').reset_index()
    instructions = pd.DataFrame({'Instructions': [
        'Run Abstract Cleaning.ipynb from top to bottom to rebuild this workbook and the analysis inputs.',
        'Only KEEP records enter the analysis. REVIEW records are withheld until explicitly accepted.',
        'This is a lexical screen. Inspect the evidence and full abstracts; decisions are not manual certification.',
        'Default scope excludes all VC mentions (including mixed PE/VC), chemistry and identified trade news.',
        'Edit configs/abstract_scope_rules.json to change screening rules.',
        'For a paper-level override copy paper_id and record_sha256 into configs/abstract_scope_overrides.csv.',
        'Set decision to keep, exclude or review and supply a reason, then rerun the notebook.',
        'Excel is a review export; edits to this workbook are not imported. Use the CSV or notebook override cell.',
        'Original text, IDs, metadata and original embeddings are preserved. Only corpus rows are selected.',
        'Next run Clustering Abstracts.ipynb, then Topic_analysis_Abstracts_v2.ipynb or Topic analysis Abstracts.ipynb.',
        'Previous cluster memberships and names describe the previous corpus; refit them for the new sample.',
    ]})
    front = ['paper_id', 'decision', 'reason', 'title', 'journal', 'abstract', 'record_sha256',
             'automatic_decision', 'automatic_reason', 'manual_override']
    ordered = audit[front + [c for c in audit if c not in front and c != 'embedding_text']]
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(destination, engine='openpyxl') as writer:
        instructions.to_excel(writer, sheet_name='Read me', index=False)
        counts.to_excel(writer, sheet_name='Summary', index=False)
        for name, frame in [('All decisions', ordered), ('Keep', ordered[ordered.decision.eq('keep')]),
                            ('Excluded', ordered[ordered.decision.eq('exclude')]),
                            ('Review', ordered[ordered.decision.eq('review')])]:
            frame.to_excel(writer, sheet_name=name, index=False)
        for sheet in writer.book:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='234E52')
                sheet.column_dimensions[cell.column_letter].width = (
                    85 if cell.value in {'Instructions', 'abstract'} else
                    65 if cell.value == 'title' else 32)
            for cells in sheet.iter_rows(min_row=2):
                for cell in cells:
                    # Treat all source strings as text, including Excel formula prefixes.
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
                sheet.row_dimensions[cells[0].row].height = 45


def build_scope_bundle(root, source_csv, parent_run, rules, overrides=None):
    """Validate the parent, screen, and write a content-addressed derivative run.

    Reuse vectors by EID only after exact source/text validation. No model is run.
    Validate the finished derivative with the same loader used by clustering.
    """
    root, source_csv, parent_run = Path(root).resolve(), Path(source_csv).resolve(), Path(parent_run).resolve()
    papers, embeddings, provenance = load_abstract_embeddings(parent_run, source_csv)
    source = pd.read_csv(source_csv, dtype=str, keep_default_na=False)
    audit = screen_abstracts(source, rules, overrides)
    kept = source.loc[audit.decision.eq('keep')].copy()
    if len(kept) < 20:
        raise ValueError('Fewer than 20 retained abstracts; review scope before running the k=2–15 workflow.')
    identity = {'schema_version': 1, 'source_sha256': file_hash(source_csv),
                'parent_hashes': provenance['input_hashes'], 'rules': rules,
                'decisions': audit[['paper_id', 'record_sha256', 'decision', 'reason']].to_dict('records'),
                'code_sha256': file_hash(Path(__file__))}
    scope_id = _json_hash(identity)[:16]
    destination = root / 'data/abstracts/private_equity' / scope_id
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        manifest = json.loads((destination / 'scope_manifest.json').read_text())
        if manifest['identity'] != identity:
            raise ValueError('Existing scope bundle identity differs.')
        for name, digest in manifest['artifacts_sha256'].items():
            if file_hash(destination / name) != digest:
                raise ValueError(f'Existing scope artifact changed: {name}')
        load_abstract_embeddings(destination / 'embedding_run', destination / 'abstracts_clean.csv')
        return destination, manifest, audit
    with tempfile.TemporaryDirectory(prefix='.scope-', dir=destination.parent) as temporary:
        staging = Path(temporary) / 'bundle'
        run = staging / 'embedding_run'
        run.mkdir(parents=True)
        clean = staging / 'abstracts_clean.csv'
        kept.to_csv(clean, index=False)
        audit.to_csv(staging / 'scope_audit.csv', index=False)
        audit.loc[audit.decision.eq('exclude')].to_csv(staging / 'excluded.csv', index=False)
        audit.loc[audit.decision.eq('review')].to_csv(staging / 'review.csv', index=False)
        audit[OVERRIDE_COLUMNS].assign(decision='', reason='').to_csv(staging / 'override_template.csv', index=False)
        _write_json(staging / 'rules.json', rules)
        retained = set(kept.paper_id)
        positions = np.flatnonzero(papers.paper_id.isin(retained).to_numpy())
        selected = pd.read_csv(parent_run / 'selected_papers.csv', dtype=str, keep_default_na=False).iloc[positions].copy()
        selected['embedding_row'] = np.arange(len(selected))
        selected.to_csv(run / 'selected_papers.csv', index=False)
        selected[['embedding_row', 'paper_id']].to_csv(run / 'paper_ids.csv', index=False)
        token_audit = pd.read_csv(parent_run / 'token_lengths.csv', dtype=str, keep_default_na=False)
        token_audit.set_index('paper_id', drop=False).loc[selected.paper_id].to_csv(run / 'token_lengths.csv', index=False)
        vectors = embeddings[positions]
        np.save(run / 'embeddings.npy', vectors)
        pd.DataFrame({'paper_id': selected.paper_id.to_numpy(), 'parent_embedding_row': positions,
                      'embedding_row': np.arange(len(selected))}).to_csv(staging / 'embedding_row_map.csv', index=False)
        parent_summary = json.loads((parent_run / 'run_summary.json').read_text())
        # Preserve model facts, but never claim that parent inference timings describe this selection.
        settings = {key: parent_summary[key] for key in (
            'model', 'model_revision', 'input_column', 'input_format', 'pooling', 'normalization',
            'dtype', 'max_tokens', 'truncation', 'chunking', 'packages') if key in parent_summary}
        summary = {**settings, 'status': 'complete', 'run_mode': 'full',
                   'source': str(destination / 'abstracts_clean.csv'), 'source_sha256': file_hash(clean),
                   'corpus_size': len(kept), 'selected_size': len(kept), 'shape': list(vectors.shape),
                   'selected_over_limit': 0, 'corpus_over_limit': 0,
                   'maximum_sample_tokens': int(selected.token_count.astype(int).max()),
                   'maximum_corpus_tokens': int(selected.token_count.astype(int).max()),
                   'derivation': 'exact subset of existing embeddings; no inference',
                   'scope_id': scope_id, 'parent_run': str(parent_run),
                   'parent_input_hashes': provenance['input_hashes']}
        _write_json(run / 'run_summary.json', summary)
        loaded, checked, _ = load_abstract_embeddings(run, clean)
        if loaded.paper_id.tolist() != selected.paper_id.tolist() or not np.array_equal(checked, vectors):
            raise ValueError('Derivative embedding roundtrip failed.')
        export_review_workbook(audit, staging / 'Abstract Cleaning.xlsx')
        manifest = {'scope_id': scope_id, 'identity': identity, 'source_rows': len(source),
                    'counts': audit.decision.value_counts().to_dict(),
                    'reason_counts': audit.groupby(['decision', 'reason']).size().rename('records').reset_index().to_dict('records'),
                    'embedding_inference_performed': False,
                    'review_policy': 'Only keep enters analysis; review is withheld.',
                    'artifacts_sha256': {str(p.relative_to(staging)): file_hash(p)
                                         for p in staging.rglob('*') if p.is_file()}}
        _write_json(staging / 'scope_manifest.json', manifest)
        staging.rename(destination)
    return destination, manifest, audit


def activate_scope(root, destination):
    """Set the shared input pointer only after complete corpus/vector validation."""
    root, destination = Path(root).resolve(), Path(destination).resolve()
    source, run = destination / 'abstracts_clean.csv', destination / 'embedding_run'
    papers, _, provenance = load_abstract_embeddings(run, source)
    manifest = json.loads((destination / 'scope_manifest.json').read_text())
    for name, digest in manifest['artifacts_sha256'].items():
        if file_hash(destination / name) != digest:
            raise ValueError(f'Scope artifact changed: {name}')
    config = {'schema_version': 1, 'source_csv': str(source.relative_to(root)),
              'run_dir': str(run.relative_to(root)), 'scope_id': manifest['scope_id'],
              'corpus_size': len(papers), 'input_hashes': provenance['input_hashes'],
              'scope_manifest': str((destination / 'scope_manifest.json').relative_to(root)),
              'scope_manifest_sha256': file_hash(destination / 'scope_manifest.json')}
    path = root / ACTIVE_CONFIG
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix('.tmp')
    _write_json(temporary, config)
    temporary.replace(path)
    shutil.copyfile(destination / 'Abstract Cleaning.xlsx', root / 'Abstract Cleaning.xlsx')
    return config


def resolve_abstract_inputs(root, *, use_original=False):
    """Resolve an explicit shared pointer; never guess from modification times."""
    root = Path(root).resolve()
    pointer = root / ACTIVE_CONFIG
    if use_original or not pointer.exists():
        return root / ORIGINAL_RUN, root / ORIGINAL_SOURCE
    config = json.loads(pointer.read_text())
    if config.get('schema_version') != 1:
        raise ValueError('Unsupported active abstract input schema.')
    run, source = root / config['run_dir'], root / config['source_csv']
    for name, digest in config['input_hashes'].items():
        path = source if name == source.name else run / name
        if not path.is_file() or file_hash(path) != digest:
            raise ValueError(f'Active abstract input changed or is missing: {path}. Rerun Abstract Cleaning.ipynb.')
    manifest = root / config['scope_manifest']
    if file_hash(manifest) != config['scope_manifest_sha256']:
        raise ValueError('Active scope manifest changed; rerun Abstract Cleaning.ipynb.')
    return run, source


def abstract_partition_folder(provenance):
    """Preserve historical six-cluster exports; new scope uses current counts."""
    return 'selected_clusters' if provenance.get('scope_id') else 'six_clusters'
